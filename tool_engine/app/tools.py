import ast
import io
import operator
import sys
from pathlib import Path
from typing import Any, Dict, Optional
import fitz
import pandas as pd
import pytesseract
import timeout_decorator
from docx import Document
from PIL import Image
from pptx import Presentation
from app.config import settings
# ---------- Helpers ----------
def _safe_result(tool_name: str, result: Any = None, output_file: Optional[str] = None, error: Optional[str] = None) -> Dict[str, Any]:
    return {
        "tool": tool_name,
        "status": "error" if error else "success",
        "result": result,
        "output_file": output_file,
        "error": error,
    }
def _validate_file(path: Path, allowed_extensions: set[str]) -> Optional[str]:
    if not path.exists():
        return f"File not found: {path}"
    if path.stat().st_size == 0:
        return "File is empty."
    if path.suffix.lower() not in allowed_extensions:
        return f"File type not allowed: {path.suffix}"
    return None
def _is_inside_sandbox(path: Path) -> bool:
    try:
        path.resolve().relative_to(settings.sandbox_dir.resolve())
        return True
    except ValueError:
        return False
# ---------- Tools ----------
class PDFReaderTool:
    name = "pdf_reader"
    def execute(self, file_path: str, max_pages: Optional[int] = None, **kwargs) -> Dict[str, Any]:
        try:
            path = Path(file_path)
            err = _validate_file(path, {".pdf"})
            if err:
                return _safe_result(self.name, error=err)
            doc = fitz.open(path)
            pages = []
            limit = max_pages if max_pages else len(doc)
            for i in range(min(limit, len(doc))):
                text = doc[i].get_text("text")
                pages.append({"page": i + 1, "text": text.strip()})
            doc.close()
            return _safe_result(
                self.name,
                result={"file": str(path), "total_pages": len(pages), "pages": pages},
            )
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class OCRTool:
    name = "ocr_tool"
    def __init__(self):
        self.tesseract_available = Path(settings.tesseract_cmd).exists()
        if self.tesseract_available:
            pytesseract.pytesseract.tesseract_cmd = str(settings.tesseract_cmd)
    def _ocr_image(self, image: Image.Image) -> tuple[str, Optional[float]]:
        if not self.tesseract_available:
            raise RuntimeError("Tesseract OCR is not installed.")
        data = pytesseract.image_to_data(image, lang="eng", output_type=pytesseract.Output.DICT)
        words = []
        confs = []
        for word, conf in zip(data["text"], data["conf"]):
            word = word.strip()
            if word:
                words.append(word)
                try:
                    c = float(conf)
                    if c >= 0:
                        confs.append(c)
                except (ValueError, TypeError):
                    pass
        avg = None
        if confs:
            avg = round(sum(confs) / len(confs), 2)
        return " ".join(words), avg
    def execute(self, file_path: str, **kwargs) -> Dict[str, Any]:
        try:
            path = Path(file_path)
            err = _validate_file(path, {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".pdf"})
            if err:
                return _safe_result(self.name, error=err)
            ext = path.suffix.lower()
            if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
                image = Image.open(path).convert("RGB")
                text, conf = self._ocr_image(image)
                return _safe_result(
                    self.name,
                    result={"file": str(path), "text": text, "ocr_confidence": conf},
                )
            elif ext == ".pdf":
                doc = fitz.open(path)
                page = doc[0]
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                doc.close()
                text, conf = self._ocr_image(image)
                return _safe_result(
                    self.name,
                    result={"file": str(path), "text": text, "ocr_confidence": conf},
                )
            else:
                return _safe_result(self.name, error="Unsupported file type for OCR.")
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class ExcelAnalyzerTool:
    name = "excel_analyzer"
    def execute(self, file_path: str, sheet_name: str | int | None = None, operation: str = "summary", **kwargs) -> Dict[str, Any]:
        try:
            path = Path(file_path)
            err = _validate_file(path, {".xlsx", ".xls", ".csv"})
            if err:
                return _safe_result(self.name, error=err)
            if path.suffix.lower() == ".csv":
                df = pd.read_csv(path)
            else:
                df = pd.read_excel(path, sheet_name=sheet_name)
            if operation == "summary":
                data = {
                    "rows": len(df),
                    "columns": list(df.columns),
                    "dtypes": {str(k): str(v) for k, v in df.dtypes.items()},
                    "head": df.head(5).to_dict(orient="records"),
                }
            elif operation == "describe":
                data = df.describe(include="all").fillna("").to_dict()
            else:
                return _safe_result(self.name, error=f"Unsupported operation: {operation}")
            return _safe_result(self.name, result={"type": operation, "data": data})
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class CalculatorTool:
    name = "calculator"
    ALLOWED_NODES = {
        ast.Expression, ast.Constant, ast.Num,
        ast.UnaryOp, ast.UAdd, ast.USub,
        ast.BinOp, ast.Add, ast.Sub, ast.Mult,
        ast.Div, ast.FloorDiv, ast.Mod, ast.Pow,
    }
    OPS = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.FloorDiv: operator.floordiv,
        ast.Mod: operator.mod,
        ast.Pow: operator.pow,
    }
    def _eval(self, node):
        if type(node) not in self.ALLOWED_NODES:
            raise ValueError("Unsafe expression.")
        if isinstance(node, ast.Constant):
            return node.value
        if isinstance(node, ast.Num):
            return node.n
        if isinstance(node, ast.UnaryOp):
            val = self._eval(node.operand)
            if isinstance(node.op, ast.UAdd):
                return +val
            if isinstance(node.op, ast.USub):
                return -val
        if isinstance(node, ast.BinOp):
            l = self._eval(node.left)
            r = self._eval(node.right)
            op = type(node.op)
            if op in self.OPS:
                return self.OPS[op](l, r)
        raise ValueError("Unsafe expression.")
    def execute(self, expression: str, **kwargs) -> Dict[str, Any]:
        try:
            tree = ast.parse(expression, mode="eval")
            value = self._eval(tree.body)
            return _safe_result(self.name, result={"expression": expression, "value": value})
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class FileIOTool:
    name = "file_io"
    def execute(self, operation: str, file_path: str, content: Optional[str] = None, **kwargs) -> Dict[str, Any]:
        try:
            path = Path(file_path)
            if operation == "read":
                if not path.exists():
                    return _safe_result(self.name, error=f"File not found: {path}")
                if path.stat().st_size == 0:
                    return _safe_result(self.name, error="File is empty.")
                text = path.read_text(encoding="utf-8", errors="ignore")
                return _safe_result(self.name, result={"file": str(path), "content": text})
            elif operation == "write":
                if not _is_inside_sandbox(path):
                    return _safe_result(self.name, error="Writes allowed only inside sandbox directory.")
                if content is None:
                    return _safe_result(self.name, error="No content provided for write.")
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(content, encoding="utf-8")
                return _safe_result(self.name, result={"file": str(path), "bytes_written": len(content)}, output_file=str(path))
            else:
                return _safe_result(self.name, error=f"Unsupported operation: {operation}")
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class WordGeneratorTool:
    name = "word_generator"
    def execute(self, title: str, sections: list[dict], output_filename: str, **kwargs) -> Dict[str, Any]:
        try:
            output_path = settings.sandbox_dir / output_filename
            if not _is_inside_sandbox(output_path):
                return _safe_result(self.name, error="Output file must be inside sandbox directory.")
            doc = Document()
            doc.add_heading(title, 0)
            for sec in sections:
                heading = sec.get("heading", "")
                body = sec.get("body", "")
                if heading:
                    doc.add_heading(heading, level=1)
                if body:
                    doc.add_paragraph(body)
            doc.save(output_path)
            return _safe_result(
                self.name,
                result={"title": title, "sections_count": len(sections)},
                output_file=str(output_path),
            )
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class PPTGeneratorTool:
    name = "ppt_generator"
    def execute(self, title: str, slides: list[dict], output_filename: str, **kwargs) -> Dict[str, Any]:
        try:
            output_path = settings.sandbox_dir / output_filename
            if not _is_inside_sandbox(output_path):
                return _safe_result(self.name, error="Output file must be inside sandbox directory.")
            prs = Presentation()
            title_slide = prs.slides.add_slide(prs.slide_layouts[0])
            title_slide.shapes.title.text = title
            for s in slides:
                heading = s.get("heading", "")
                content = s.get("content", "")
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = heading
                slide.placeholders[1].text = content
            prs.save(output_path)
            return _safe_result(
                self.name,
                result={"title": title, "slides_count": len(slides)},
                output_file=str(output_path),
            )
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class CodeSandboxTool:
    name = "code_sandbox"
    def _run(self, code: str, timeout: int):
        @timeout_decorator.timeout(timeout)
        def run():
            allowed_builtins = {
                "print": print,
                "len": len,
                "sum": sum,
                "min": min,
                "max": max,
                "range": range,
                "str": str,
                "int": int,
                "float": float,
                "list": list,
                "dict": dict,
            }
            old_stdout = sys.stdout
            sys.stdout = buf = io.StringIO()
            try:
                exec(code, {"__builtins__": allowed_builtins}, {})
                return buf.getvalue()
            finally:
                sys.stdout = old_stdout
        return run()
    def execute(self, code: str, **kwargs) -> Dict[str, Any]:
        try:
            output = self._run(code, settings.code_timeout)
            return _safe_result(self.name, result={"code": code, "output": output})
        except Exception as e:
            return _safe_result(self.name, error=str(e))
class ReportGeneratorTool:
    name = "report_generator"
    def execute(
        self,
        report_title: str,
        summary_text: str,
        word_filename: str,
        ppt_filename: str,
        **kwargs,
    ) -> Dict[str, Any]:
        try:
            word_tool = WordGeneratorTool()
            ppt_tool = PPTGeneratorTool()
            sections = [
                {"heading": "Summary", "body": summary_text},
            ]
            slides = [
                {"heading": "Overview", "content": summary_text[:300]},
            ]
            word_res = word_tool.execute(
                title=report_title,
                sections=sections,
                output_filename=word_filename,
            )
            if word_res["status"] == "error":
                return _safe_result(self.name, error=word_res["error"])
            ppt_res = ppt_tool.execute(
                title=report_title,
                slides=slides,
                output_filename=ppt_filename,
            )
            if ppt_res["status"] == "error":
                return _safe_result(self.name, error=ppt_res["error"])
            return _safe_result(
                self.name,
                result={
                    "report_title": report_title,
                    "word_file": word_res["output_file"],
                    "ppt_file": ppt_res["output_file"],
                },
                output_file=word_res["output_file"],
            )
        except Exception as e:
            return _safe_result(self.name, error=str(e))